import pytest
from docx import Document
from fpdf import FPDF
from tempfile import NamedTemporaryFile
import os

FILE_CONTENT = "This is sample document"

@pytest.fixture
def generate_pdf_file():
    with NamedTemporaryFile(suffix=".pdf",mode="rb",delete=False) as temp_file:
        pdf = FPDF()
        pdf.add_page()
        pdf.cell(200,100,FILE_CONTENT)
        pdf.output(temp_file.name)
    yield temp_file.name
    os.remove(temp_file.name)

@pytest.fixture
def generate_doc_file():
    with NamedTemporaryFile(suffix=".docx",mode="rb",delete=False) as temp_file:
       doc = Document()
       doc.add_paragraph(FILE_CONTENT)
       doc.save(temp_file.name)
    yield temp_file.name
    os.remove(temp_file.name)